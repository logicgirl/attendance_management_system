"""Attendance reports as formatted Excel workbooks."""

import datetime as dt
import io
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from app.models import Attendance

HEADER_ROW = 4  # rows 1-2 hold the title and the filters used

STATUS_STYLE = {
    "Present": ("E8F5E9", "1B5E20"),
    "Late": ("FFF8E1", "8D6E00"),
    "Absent": ("FDECEA", "B71C1C"),
}
LOW_RATE = 0.75


def _title(ws: Worksheet, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color="1F2937")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color="6B7280")


def _table(ws: Worksheet, name: str, headers: list[str], rows: list[list], widths: list[int]):
    for col, (header, width) in enumerate(zip(headers, widths, strict=True), start=1):
        ws.cell(row=HEADER_ROW, column=col, value=header)
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(rows, start=HEADER_ROW + 1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    last_col = get_column_letter(len(headers))
    if rows:
        # An Excel table gives banded rows and a filter button on every column
        table = Table(displayName=name, ref=f"A{HEADER_ROW}:{last_col}{HEADER_ROW + len(rows)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    else:
        for col in range(1, len(headers) + 1):
            ws.cell(row=HEADER_ROW, column=col).font = Font(bold=True)
        ws.cell(row=HEADER_ROW + 1, column=1, value="No records match these filters.")
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)


def attendance_workbook(
    records: Sequence[Attendance], stats: list[dict], filters_text: str, generated: dt.datetime
) -> bytes:
    wb = Workbook()
    subtitle = f"{filters_text}  ·  Generated {generated:%d %b %Y, %H:%M}"

    ws = wb.active
    ws.title = "Attendance"
    _title(ws, "AttendEase attendance report", subtitle)
    headers = [
        "Date",
        "Day",
        "Time",
        "Student ID",
        "Full Name",
        "Course",
        "Status",
        "Term",
        "Record ID",
    ]
    rows = [
        [
            r.date,
            r.day,
            r.time,
            r.student_id,
            r.student_name,
            r.course_code,
            r.status.value,
            r.term,
            r.id,
        ]
        for r in records
    ]
    _table(ws, "Attendance", headers, rows, [13, 12, 10, 16, 26, 12, 11, 14, 11])
    for i, row in enumerate(rows, start=HEADER_ROW + 1):
        ws.cell(row=i, column=1).number_format = "dd mmm yyyy"
        ws.cell(row=i, column=3).number_format = "h:mm AM/PM"
        fill, color = STATUS_STYLE.get(row[6], ("FFFFFF", "000000"))
        status = ws.cell(row=i, column=7)
        status.fill = PatternFill("solid", fgColor=fill)
        status.font = Font(bold=True, color=color)
        status.alignment = Alignment(horizontal="center")

    ws2 = wb.create_sheet("By student")
    _title(
        ws2,
        "Attendance by student",
        f"Late counts as attended · below 75% is highlighted · {subtitle}",
    )
    headers2 = [
        "Student ID",
        "Full Name",
        "Sessions",
        "Present",
        "Late",
        "Absent",
        "Attendance rate",
    ]
    rows2 = [
        [
            s["student_id"],
            s["name"],
            s["total"],
            s["present"],
            s["late"],
            s["absent"],
            s["attendance_rate"] / 100,
        ]
        for s in stats
    ]
    _table(ws2, "ByStudent", headers2, rows2, [16, 26, 11, 10, 9, 10, 16])
    for i, row in enumerate(rows2, start=HEADER_ROW + 1):
        rate = ws2.cell(row=i, column=7)
        rate.number_format = "0.0%"
        if row[6] < LOW_RATE:
            rate.fill = PatternFill("solid", fgColor=STATUS_STYLE["Absent"][0])
            rate.font = Font(bold=True, color=STATUS_STYLE["Absent"][1])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
