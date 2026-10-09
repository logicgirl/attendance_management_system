import datetime as dt
import io

import pytest
from openpyxl import load_workbook

from app.services.attendance import now_local

pytestmark = pytest.mark.anyio


@pytest.fixture
async def course_id(make_student, make_course, lecturer):
    ids = tuple(f"CS-00{i}" for i in range(1, 5))
    for i, student_id in enumerate(ids):
        await make_student(student_id, name=f"Student {'ABCD'[i]}")
    return await make_course("CS101", lecturer.id, ids)


async def register(client, headers, course_id, statuses: dict[str, str]):
    res = await client.post(
        "/api/v1/attendance/register",
        headers=headers,
        json={
            "course_id": course_id,
            "records": [{"student_id": s, "status": v} for s, v in statuses.items()],
        },
    )
    assert res.status_code == 200, res.text


async def test_summary(client, lecturer_headers, course_id):
    await register(
        client,
        lecturer_headers,
        course_id,
        {"CS-001": "Present", "CS-002": "Present", "CS-003": "Late"},
    )
    res = await client.get(
        f"/api/v1/reports/summary?course_id={course_id}", headers=lecturer_headers
    )
    body = res.json()
    assert body["active_students"] == 4
    assert body["total_records"] == 3
    assert (body["present"], body["late"], body["absent"]) == (2, 1, 0)
    assert body["present_rate"] == 66.7


async def test_lecturer_summary_only_counts_own_students(
    client, lecturer_headers, course_id, make_student
):
    await make_student("OUTSIDER")
    body = (await client.get("/api/v1/reports/summary", headers=lecturer_headers)).json()
    assert body["active_students"] == 4
    assert body["total_records"] == 0
    assert body["present_rate"] == 0.0


async def test_student_stats(client, lecturer_headers, course_id):
    await register(client, lecturer_headers, course_id, {"CS-001": "Late", "CS-002": "Absent"})
    res = await client.get(
        f"/api/v1/reports/students?course_id={course_id}", headers=lecturer_headers
    )
    stats = {s["student_id"]: s for s in res.json()}
    assert stats["CS-001"]["attendance_rate"] == 100.0
    assert stats["CS-002"]["attendance_rate"] == 0.0
    assert stats["CS-002"]["absent"] == 1


async def test_excel_export(client, lecturer_headers, course_id):
    await register(client, lecturer_headers, course_id, {"CS-001": "Present", "CS-002": "Absent"})

    res = await client.get(
        f"/api/v1/reports/attendance.xlsx?course_id={course_id}", headers=lecturer_headers
    )
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert ".xlsx" in res.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Attendance", "By student"]
    sheet = wb["Attendance"]
    assert "Course CS101" in sheet["A2"].value
    header = [c.value for c in sheet[4]]
    assert header[:7] == ["Date", "Day", "Time", "Student ID", "Full Name", "Course", "Status"]
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=5)]
    assert {(r[3], r[6]) for r in rows} == {("CS-001", "Present"), ("CS-002", "Absent")}
    assert rows[0][0].date() == now_local().date()  # a real date, not text

    by_student = {r[0].value: r for r in wb["By student"].iter_rows(min_row=5)}
    assert by_student["CS-001"][6].value == 1.0
    assert by_student["CS-001"][6].number_format == "0.0%"
    assert by_student["CS-002"][6].value == 0.0


async def test_excel_export_with_no_records(client, lecturer_headers, course_id):
    res = await client.get("/api/v1/reports/attendance.xlsx", headers=lecturer_headers)
    sheet = load_workbook(io.BytesIO(res.content))["Attendance"]
    assert sheet["A5"].value == "No records match these filters."


async def test_daily_counts(client, lecturer_headers, course_id):
    today = now_local().date()
    yesterday = (today - dt.timedelta(days=1)).isoformat()
    res = await client.post(
        "/api/v1/attendance/register",
        headers=lecturer_headers,
        json={
            "course_id": course_id,
            "date": yesterday,
            "records": [
                {"student_id": "CS-001", "status": "Present"},
                {"student_id": "CS-002", "status": "Absent"},
            ],
        },
    )
    assert res.status_code == 200
    await register(client, lecturer_headers, course_id, {"CS-001": "Late"})

    days = (await client.get("/api/v1/reports/daily?days=3", headers=lecturer_headers)).json()
    assert [d["date"] for d in days][-1] == today.isoformat()
    assert len(days) == 3
    assert days[0]["total"] == 0 and days[0]["attendance_rate"] is None
    assert (days[1]["present"], days[1]["absent"], days[1]["attendance_rate"]) == (1, 1, 50.0)
    assert (days[2]["late"], days[2]["attendance_rate"]) == (1, 100.0)


async def test_course_overview(client, lecturer_headers, course_id, make_course):
    await make_course("OTHER")  # not taught by this lecturer
    await register(client, lecturer_headers, course_id, {"CS-001": "Present", "CS-002": "Late"})

    overview = (await client.get("/api/v1/reports/courses", headers=lecturer_headers)).json()
    assert len(overview) == 1
    row = overview[0]
    assert (row["code"], row["enrolled"], row["marked"]) == ("CS101", 4, 2)
    assert (row["present"], row["late"], row["absent"]) == (1, 1, 0)
    assert row["lecturer_name"] == "Lecturer"
